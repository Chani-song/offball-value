"""Exercise persistence on temporary copies, never the canonical labels."""

import csv
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "tools/shot_annotation_app/app.py"
CANONICAL = ROOT / "annotations/shot_annotations.xlsx"


def load_app():
    spec = importlib.util.spec_from_file_location("shot_annotation_app", APP_PATH)
    module = importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules, {spec.name: module}):
        spec.loader.exec_module(module)
    return module


class AnnotationAppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.media = self.root / "media"
        self.media.mkdir()
        self.xlsx = self.root / "labels.xlsx"
        shutil.copyfile(CANONICAL, self.xlsx)
        self.env = patch.dict(os.environ, {
            "OFFBALL_ANNOTATION_MEDIA_DIR": str(self.media),
            "OFFBALL_ANNOTATION_XLSX": str(self.xlsx),
        })
        self.env.start()
        self.addCleanup(self.env.stop)
        self.module = load_app()
        self.original = self.module._annotation_rows()
        self.row = next(iter(self.original.values()))
        self.clip = self.media / self.row["video_filename"]
        self.clip.write_bytes(b"test clip")
        self.client = self.module.app.test_client()

    def test_startup_reads_canonical_default_without_writing(self):
        before = CANONICAL.read_bytes()
        with patch.dict(os.environ):
            os.environ.pop("OFFBALL_ANNOTATION_XLSX", None)
            module = load_app()
            self.assertEqual(module.ANNOT_XLSX, CANONICAL)
            client = module.app.test_client()
            self.assertEqual(client.get("/").status_code, 200)
            data = client.get("/api/shots").get_json()
            self.assertTrue(data["xlsx_exists"])
            self.assertEqual(data["shots"][0]["annotation"], self.row)
            self.assertEqual(data["reviewed_count"], 1)
        self.assertEqual(CANONICAL.read_bytes(), before)

    def test_save_updates_one_row_and_survives_process_restart(self):
        clip_id = self.row["clip_id"]
        for effect in ("strong", "medium", "low", "ignore"):
            with self.subTest(effect=effect):
                payload = {
                    "clip_id": clip_id,
                    "offball_attackers": ["7", "11"],
                    "drawn_defenders": ["2", "4"],
                    "space_beneficiaries": ["9", "27"],
                    "effect": effect,
                    "notes": "Temporary test: comments · 메모",
                }
                response = self.client.post("/api/save", json=payload)
                self.assertEqual(response.status_code, 200)
                saved = response.get_json()["row"]
                self.assertEqual(saved["offball_attackers"], "7, 11")
                self.assertEqual(saved["drawn_defenders"], "2, 4")
                self.assertEqual(saved["space_beneficiaries"], "9, 27")
                self.assertEqual(saved["team"], self.row["team"])
                # A fresh interpreter must recover every row without CSV.
                self.module.ANNOT_CSV.unlink()
                code = (
                    "import json, runpy; "
                    f"m = runpy.run_path({str(APP_PATH)!r}); "
                    "print(json.dumps(m['_annotation_rows']()))"
                )
                restored = json.loads(subprocess.check_output([sys.executable, "-c", code], text=True))
                self.assertEqual(restored.pop(clip_id), saved)
                self.assertEqual(restored, {k: v for k, v in self.original.items() if k != clip_id})
                wb = load_workbook(self.xlsx, read_only=True)
                try:
                    self.assertEqual(wb.active.max_row, len(self.original) + 1)
                finally:
                    wb.close()

    def test_csv_export_uses_xlsx_and_does_not_rewrite_it(self):
        before = self.xlsx.read_bytes()
        self.module.ANNOT_CSV.write_text("clip_id,notes\nstale,wrong\n")
        self.assertEqual(self.module._annotation_rows(), self.original)
        response = self.client.get("/download/csv")
        self.assertEqual(response.status_code, 200)
        response.close()
        with self.module.ANNOT_CSV.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual({r["clip_id"]: r for r in rows}, self.original)
        self.assertEqual(self.xlsx.read_bytes(), before)

    def test_clip_layouts_keep_the_same_id_and_serve_video(self):
        for layout in ("direct", "collected", "per_match"):
            with self.subTest(layout=layout):
                media = self.root / layout
                if layout == "collected":
                    clip = media / "ALL_SHOT_CLIPS" / self.row["video_filename"]
                elif layout == "per_match":
                    clip = media / self.row["match_id"] / "clips" / (self.row["clip_id"].split(":")[1] + ".mp4")
                else:
                    clip = media / self.row["video_filename"]
                clip.parent.mkdir(parents=True)
                clip.write_bytes(b"test clip")
                if layout == "per_match":
                    # The legacy output also contains a combined reel.
                    (media / "ALL_7_MATCHES_SHOTS.mp4").write_bytes(b"reel")
                with patch.dict(os.environ, OFFBALL_ANNOTATION_MEDIA_DIR=str(media)):
                    client = load_app().app.test_client()
                    shots = client.get("/api/shots").get_json()["shots"]
                    self.assertEqual(len(shots), 1)
                    self.assertEqual(shots[0]["clip_id"], self.row["clip_id"])
                    response = client.get("/video/" + self.row["clip_id"])
                    self.assertEqual(response.data, b"test clip")
                    response.close()

    def test_missing_media_prints_configuration_help_at_startup(self):
        with patch.dict(os.environ, OFFBALL_ANNOTATION_MEDIA_DIR=str(self.root / "missing")):
            code = (
                "import runpy; from unittest.mock import patch; "
                "from flask import Flask; "
                "mock = patch.object(Flask, 'run'); mock.start(); "
                f"runpy.run_path({str(APP_PATH)!r}, run_name='__main__')"
            )
            output = subprocess.check_output([sys.executable, "-c", code], text=True)
        self.assertIn("Media directory unavailable", output)
        self.assertIn("export OFFBALL_ANNOTATION_MEDIA_DIR=/path/to/shot/clips", output)
        self.assertIn("Clips found:  0", output)


if __name__ == "__main__":
    unittest.main()
