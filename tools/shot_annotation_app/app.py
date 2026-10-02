from __future__ import annotations

import bisect
import csv
import json
import os
import re
import socket
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIR.parents[1]
OUTPUT_DIR = Path(os.environ.get(
    "OFFBALL_ANNOTATION_MEDIA_DIR", Path.home() / "idsse_shots" / "output"
)).expanduser().resolve()
COLLECTED_DIR = OUTPUT_DIR / "ALL_SHOT_CLIPS"
if not COLLECTED_DIR.is_dir():
    COLLECTED_DIR = OUTPUT_DIR
META_PRIMARY = COLLECTED_DIR / "shot_metadata.csv"
META_FALLBACK = OUTPUT_DIR / "ALL_7_MATCHES_SHOTS.csv"
ANNOT_XLSX = Path(os.environ.get(
    "OFFBALL_ANNOTATION_XLSX", PROJECT_ROOT / "annotations" / "shot_annotations.xlsx"
)).expanduser().resolve()
ANNOT_CSV = ANNOT_XLSX.with_suffix(".csv")
MEDIA_ROOT = OUTPUT_DIR.parent if OUTPUT_DIR.name == "ALL_SHOT_CLIPS" else OUTPUT_DIR
DATA_DIR = Path(os.environ.get("IDSSE_ROOT", MEDIA_ROOT.parent)).expanduser().resolve() / "idsse-data"
TRACKING_CACHE_DIR = OUTPUT_DIR / "tracking_overlay_cache"
TRACKING_FPS = 5.0
TRACKING_SOURCE_FPS = 25.0
CLIP_BEFORE = 8.0
CLIP_AFTER = 3.0

_TRACKING_MEMORY: dict[str, Any] = {"match_id": None, "data": None}
_TRACKING_LOCK = threading.Lock()

MATCH_ORDER = ["J03WMX", "J03WN1", "J03WOH", "J03WOY", "J03WPY", "J03WQQ", "J03WR9"]
EFFECTS = {"strong", "medium", "low", "ignore"}

FIELDS = [
    "clip_id",
    "match_id",
    "shot_number",
    "period",
    "period_seconds",
    "match_clock",
    "team",
    "shooter",
    "shot_result",
    "video_filename",
    "offball_attackers",
    "drawn_defenders",
    "space_beneficiaries",
    "effect",
    "notes",
    "reviewed",
    "saved_at",
]

app = Flask(__name__)


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _metadata_index() -> dict[tuple[str, str], dict[str, str]]:
    meta_path = META_PRIMARY if META_PRIMARY.exists() else META_FALLBACK
    rows = _read_csv_rows(meta_path)
    idx: dict[tuple[str, str], dict[str, str]] = {}
    for row in rows:
        match_id = (row.get("match_id") or "").strip()
        clip_field = (row.get("clip") or row.get("video_filename") or "").strip()
        clip_name = Path(clip_field).name if clip_field else ""
        if match_id and clip_name:
            idx[(match_id, clip_name)] = row
    return idx


def _scan_clips() -> list[dict[str, Any]]:
    metadata = _metadata_index()
    annotations = _annotation_rows()
    clips: list[dict[str, Any]] = []

    candidates: list[tuple[str, Path]] = []
    if COLLECTED_DIR != OUTPUT_DIR or any(COLLECTED_DIR.glob("J03*_shot_*.mp4")):
        for p in sorted(COLLECTED_DIR.glob("*.mp4")):
            m = re.match(r"^(J03[A-Z0-9]+)_(shot_.*\.mp4)$", p.name, re.IGNORECASE)
            if m:
                candidates.append((m.group(1), p))
            else:
                candidates.append(("UNKNOWN", p))
    else:
        for match_dir in OUTPUT_DIR.iterdir() if OUTPUT_DIR.is_dir() else []:
            clip_dir = match_dir / "clips"
            if clip_dir.is_dir():
                for p in sorted(clip_dir.glob("*.mp4")):
                    candidates.append((match_dir.name, p))

    for match_id, path in candidates:
        original_name = path.name
        lookup_name = original_name
        if match_id != "UNKNOWN" and original_name.startswith(match_id + "_"):
            lookup_name = original_name[len(match_id) + 1 :]
        clip_id = f"{match_id}:{Path(lookup_name).stem}"
        row = {**annotations.get(clip_id, {}), **metadata.get((match_id, lookup_name), {})}

        shot_num = row.get("shot_number", "")
        if not shot_num:
            m = re.search(r"shot_(\d+)", original_name)
            shot_num = m.group(1) if m else ""

        clips.append(
            {
                "clip_id": clip_id,
                "match_id": match_id,
                "shot_number": shot_num,
                "period": row.get("period", ""),
                "period_seconds": row.get("period_seconds", ""),
                "match_clock": row.get("match_clock", ""),
                "team": row.get("team", ""),
                "shooter": row.get("player", row.get("shooter", "")),
                "shot_result": row.get("result", row.get("shot_result", "")),
                "video_filename": original_name,
                "_path": str(path.resolve()),
            }
        )

    order = {m: i for i, m in enumerate(MATCH_ORDER)}

    def key(x: dict[str, Any]):
        try:
            sn = int(x.get("shot_number") or 9999)
        except ValueError:
            sn = 9999
        return (order.get(x["match_id"], 999), x["match_id"], sn, x["video_filename"])

    clips.sort(key=key)
    return clips


def _annotation_rows() -> dict[str, dict[str, str]]:
    if not ANNOT_XLSX.exists():
        return {}
    wb = load_workbook(ANNOT_XLSX, read_only=True)
    try:
        values = wb["Shot annotations"].iter_rows(values_only=True)
        headers = next(values)
        rows = [
            {key: "" if value is None else str(value) for key, value in zip(headers, row) if key}
            for row in values
        ]
    finally:
        wb.close()
    return {r.get("clip_id", ""): r for r in rows if r.get("clip_id")}


def _normalize_numbers(value: Any) -> str:
    if isinstance(value, list):
        parts = value
    else:
        parts = re.split(r"[,;\s]+", str(value or ""))
    clean: list[str] = []
    seen = set()
    for part in parts:
        s = str(part).strip().lstrip("#")
        if not s:
            continue
        if not re.fullmatch(r"\d{1,3}", s):
            raise ValueError(f"Invalid jersey number: {part}")
        n = str(int(s))
        if n not in seen:
            seen.add(n)
            clean.append(n)
    return ", ".join(clean)


def _write_csv(rows: list[dict[str, str]]) -> None:
    ANNOT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with ANNOT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows([{k: r.get(k, "") for k in FIELDS} for r in rows])


def _write_outputs(rows_by_id: dict[str, dict[str, str]]) -> None:
    ANNOT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows_by_id.values())
    order = {m: i for i, m in enumerate(MATCH_ORDER)}

    def key(r: dict[str, str]):
        try:
            sn = int(r.get("shot_number") or 9999)
        except ValueError:
            sn = 9999
        return (order.get(r.get("match_id", ""), 999), r.get("match_id", ""), sn, r.get("video_filename", ""))

    rows.sort(key=key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Shot annotations"
    ws.append(FIELDS)
    for row in rows:
        ws.append([row.get(k, "") for k in FIELDS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "clip_id": 34,
        "match_id": 12,
        "shot_number": 12,
        "period": 9,
        "period_seconds": 15,
        "match_clock": 13,
        "team": 24,
        "shooter": 22,
        "shot_result": 16,
        "video_filename": 42,
        "offball_attackers": 22,
        "drawn_defenders": 22,
        "space_beneficiaries": 24,
        "effect": 14,
        "notes": 50,
        "reviewed": 11,
        "saved_at": 24,
    }
    for i, field in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(field, 16)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(ANNOT_XLSX)
    _write_csv(rows)



def _norm_team_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _data_file(match_id: str, kind: str) -> Path:
    patterns = {
        "meta": f"*matchinformation*{match_id}.xml",
        "tracking": f"*positions_raw_observed*{match_id}.xml",
    }
    files = list(DATA_DIR.glob(patterns[kind]))
    if len(files) != 1:
        raise FileNotFoundError(
            f"Expected exactly one {kind} file for {match_id} in {DATA_DIR}; found {len(files)}"
        )
    return files[0]


def _player_team_id(player: Any) -> str:
    try:
        return str(player.team.team_id)
    except Exception:
        return ""


def _player_jersey(player: Any) -> str:
    try:
        value = getattr(player, "jersey_no", None)
        if value is None:
            return ""
        return str(int(value))
    except Exception:
        return ""


def _load_match_tracking(match_id: str) -> dict[str, Any]:
    """Load one match at 5 Hz and keep only one match resident in memory."""
    with _TRACKING_LOCK:
        if _TRACKING_MEMORY.get("match_id") == match_id and _TRACKING_MEMORY.get("data") is not None:
            return _TRACKING_MEMORY["data"]

        try:
            from kloppy import sportec
        except Exception as e:
            raise RuntimeError(
                "kloppy is required for tracking rings. Restart with ./start.sh so requirements are installed."
            ) from e

        meta_file = _data_file(match_id, "meta")
        tracking_file = _data_file(match_id, "tracking")

        tracking = sportec.load_tracking(
            meta_data=str(meta_file),
            raw_data=str(tracking_file),
            sample_rate=TRACKING_FPS / TRACKING_SOURCE_FPS,
            only_alive=False,
        )

        teams = list(tracking.metadata.teams)
        team_info = [
            {
                "id": str(t.team_id),
                "name": str(t.name),
                "norm": _norm_team_name(t.name),
            }
            for t in teams
        ]

        frames_by_period: dict[int, list[Any]] = {}
        times_by_period: dict[int, list[float]] = {}
        for frame in tracking.frames:
            pid = int(frame.period.id)
            frames_by_period.setdefault(pid, []).append(frame)
            times_by_period.setdefault(pid, []).append(frame.timestamp.total_seconds())

        data = {
            "tracking": tracking,
            "team_info": team_info,
            "frames_by_period": frames_by_period,
            "times_by_period": times_by_period,
        }
        _TRACKING_MEMORY["match_id"] = match_id
        _TRACKING_MEMORY["data"] = data
        return data


def _tracking_cache_path(clip: dict[str, Any]) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(str(clip["video_filename"])).stem)
    return TRACKING_CACHE_DIR / str(clip["match_id"]) / f"{safe}.json"


def _build_clip_tracking(clip: dict[str, Any]) -> dict[str, Any]:
    cache_path = _tracking_cache_path(clip)
    if cache_path.exists():
        try:
            with cache_path.open("r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass

    try:
        pid = int(float(str(clip.get("period") or "0")))
        shot_sec = float(str(clip.get("period_seconds") or "nan"))
    except ValueError as e:
        raise RuntimeError("Missing period/period_seconds metadata for this clip") from e
    if pid <= 0 or shot_sec != shot_sec:
        raise RuntimeError("Missing period/period_seconds metadata for this clip")

    match_id = str(clip["match_id"])
    match = _load_match_tracking(match_id)
    period_frames = match["frames_by_period"].get(pid, [])
    period_times = match["times_by_period"].get(pid, [])
    if not period_frames:
        raise RuntimeError(f"No tracking frames found for period {pid}")

    start_sec = max(0.0, shot_sec - CLIP_BEFORE)
    end_sec = shot_sec + CLIP_AFTER
    i0 = bisect.bisect_left(period_times, start_sec)
    i1 = bisect.bisect_right(period_times, end_sec)
    clip_frames = period_frames[i0:i1]
    if not clip_frames:
        raise RuntimeError("No tracking frames found for this clip window")

    team_info = match["team_info"]
    shot_team_norm = _norm_team_name(clip.get("team", ""))
    attacking = next((t for t in team_info if t["norm"] == shot_team_norm), None)
    if attacking is None and shot_team_norm:
        attacking = next(
            (t for t in team_info if shot_team_norm in t["norm"] or t["norm"] in shot_team_norm),
            None,
        )
    if attacking is None:
        attacking = team_info[0] if team_info else {"id": "", "name": str(clip.get("team", ""))}
    defending = next((t for t in team_info if t["id"] != attacking["id"]), None)
    if defending is None:
        defending = {"id": "", "name": "Opponent"}

    attack_id = attacking["id"]
    defend_id = defending["id"]
    frames_out: list[dict[str, Any]] = []
    attack_roster: dict[str, str] = {}
    defend_roster: dict[str, str] = {}

    for frame in clip_frames:
        a: dict[str, list[float]] = {}
        d: dict[str, list[float]] = {}
        for player, point in frame.players_coordinates.items():
            if point is None:
                continue
            jersey = _player_jersey(player)
            if not jersey:
                continue
            tid = _player_team_id(player)
            try:
                name = str(getattr(player, "name", "") or "")
            except Exception:
                name = ""
            xy = [round(float(point.x), 5), round(float(point.y), 5)]
            if tid == attack_id:
                a[jersey] = xy
                if name:
                    attack_roster[jersey] = name
            elif tid == defend_id:
                d[jersey] = xy
                if name:
                    defend_roster[jersey] = name
        frames_out.append({"a": a, "d": d})

    out = {
        "ok": True,
        "clip_id": clip["clip_id"],
        "fps": TRACKING_FPS,
        "frame_count": len(frames_out),
        "attacking_team": attacking["name"],
        "defending_team": defending["name"],
        "attack_roster": attack_roster,
        "defend_roster": defend_roster,
        "frames": frames_out,
    }

    cache_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache_path.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    tmp.replace(cache_path)
    return out

def _find_clip(clip_id: str) -> dict[str, Any] | None:
    for clip in _scan_clips():
        if clip["clip_id"] == clip_id:
            return clip
    return None


@app.route("/")
def index():
    return render_template("index.html")


@app.get("/api/shots")
def api_shots():
    clips = _scan_clips()
    annotations = _annotation_rows()
    public = []
    for clip in clips:
        c = {k: v for k, v in clip.items() if not k.startswith("_")}
        c["annotation"] = annotations.get(c["clip_id"], {})
        c["reviewed"] = c["clip_id"] in annotations
        public.append(c)
    return jsonify(
        {
            "shots": public,
            "count": len(public),
            "reviewed_count": sum(1 for x in public if x["reviewed"]),
            "xlsx_exists": ANNOT_XLSX.exists(),
        }
    )


@app.get("/video/<path:clip_id>")
def video(clip_id: str):
    clip = _find_clip(clip_id)
    if clip is None:
        abort(404)
    return send_file(clip["_path"], mimetype="video/mp4", conditional=True)



@app.get("/api/tracking/<path:clip_id>")
def api_tracking(clip_id: str):
    clip = _find_clip(clip_id)
    if clip is None:
        return jsonify({"ok": False, "error": "Unknown clip_id"}), 404
    try:
        return jsonify(_build_clip_tracking(clip))
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app.post("/api/save")
def api_save():
    data = request.get_json(force=True) or {}
    clip_id = str(data.get("clip_id") or "").strip()
    clip = _find_clip(clip_id)
    if clip is None:
        return jsonify({"ok": False, "error": "Unknown clip_id"}), 404

    effect = str(data.get("effect") or "").strip().lower()
    if effect not in EFFECTS:
        return jsonify({"ok": False, "error": "Choose an effect: strong, medium, low, or ignore."}), 400

    try:
        offball = _normalize_numbers(data.get("offball_attackers", []))
        defenders = _normalize_numbers(data.get("drawn_defenders", []))
        beneficiaries = _normalize_numbers(data.get("space_beneficiaries", []))
    except ValueError as e:
        return jsonify({"ok": False, "error": str(e)}), 400

    if effect != "ignore" and (not offball or not defenders or not beneficiaries):
        return jsonify({"ok": False, "error": "For strong/medium/low, enter at least one jersey number in all three role groups."}), 400

    existing = _annotation_rows()
    row = {
        "clip_id": clip["clip_id"],
        "match_id": clip["match_id"],
        "shot_number": str(clip.get("shot_number", "")),
        "period": str(clip.get("period", "")),
        "period_seconds": str(clip.get("period_seconds", "")),
        "match_clock": str(clip.get("match_clock", "")),
        "team": str(clip.get("team", "")),
        "shooter": str(clip.get("shooter", "")),
        "shot_result": str(clip.get("shot_result", "")),
        "video_filename": str(clip.get("video_filename", "")),
        "offball_attackers": offball,
        "drawn_defenders": defenders,
        "space_beneficiaries": beneficiaries,
        "effect": effect,
        "notes": str(data.get("notes") or "").strip(),
        "reviewed": "yes",
        "saved_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
    }
    existing[clip_id] = row
    _write_outputs(existing)
    return jsonify({"ok": True, "row": row, "xlsx": str(ANNOT_XLSX)})


@app.get("/download/excel")
def download_excel():
    if not ANNOT_XLSX.exists():
        _write_outputs(_annotation_rows())
    return send_file(ANNOT_XLSX, as_attachment=True, download_name="shot_annotations.xlsx")


@app.get("/download/csv")
def download_csv():
    _write_csv(list(_annotation_rows().values()))
    return send_file(ANNOT_CSV, as_attachment=True, download_name="shot_annotations.csv")


@app.get("/api/status")
def status():
    clips = _scan_clips()
    return jsonify(
        {
            "project_root": str(PROJECT_ROOT),
            "clip_count": len(clips),
            "annotation_xlsx": str(ANNOT_XLSX),
            "annotation_csv": str(ANNOT_CSV),
        }
    )


def local_ip() -> str:
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "<MAC-IP>"


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8787"))
    clips = _scan_clips()
    print("\nIDSSE Shot Annotation")
    print(f"Project root: {PROJECT_ROOT}")
    print(f"Media folder: {OUTPUT_DIR}")
    if not OUTPUT_DIR.is_dir():
        print("Media directory unavailable. Set OFFBALL_ANNOTATION_MEDIA_DIR to your shot clips directory:")
        print("  export OFFBALL_ANNOTATION_MEDIA_DIR=/path/to/shot/clips")
    print(f"Clips found:  {len(clips)}")
    print(f"Mac browser:  http://127.0.0.1:{port}")
    print(f"Galaxy Tab:   http://{local_ip()}:{port}   (same Wi-Fi)")
    print(f"Excel output: {ANNOT_XLSX}\n")
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
